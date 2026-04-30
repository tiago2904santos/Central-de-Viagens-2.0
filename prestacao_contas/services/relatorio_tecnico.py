from __future__ import annotations

import re
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.utils import timezone
from django.utils.text import slugify

from cadastros.models import ConfiguracaoSistema
from eventos.services.documentos.renderer import (
    DocumentRendererUnavailable,
    DocumentValidationError,
    convert_docx_bytes_to_pdf_bytes,
    extract_placeholders_from_doc,
    render_docx_template_bytes,
)

from ..models import PrestacaoConta, RelatorioTecnicoPrestacao

PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]+\}\}")
RT_TEMPLATE_PATH = Path(settings.BASE_DIR) / "templates_docx" / "prestacao_contas" / "relatorio-tecnico-model.docx"
INFO_COMPLEMENTARES_FALLBACK = "(SE TROCAR DE VIATURA, FAVOR INFORMAR AQUI)"


def _formata_data_extenso_pt(dt: date) -> str:
    meses = [
        "janeiro",
        "fevereiro",
        "marco",
        "abril",
        "maio",
        "junho",
        "julho",
        "agosto",
        "setembro",
        "outubro",
        "novembro",
        "dezembro",
    ]
    return f"{dt.day} de {meses[dt.month - 1]} de {dt.year}"


def _valor_padrao(value, fallback):
    value = (str(value or "")).strip()
    return value if value else fallback


def _parse_decimal(value):
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    raw = str(value).strip()
    if not raw:
        return None
    normalized = raw.replace("R$", "").replace(" ", "")
    if "," in normalized and "." in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    else:
        normalized = normalized.replace(",", ".")
    try:
        return Decimal(normalized)
    except (InvalidOperation, TypeError, ValueError):
        return None


def _formatar_moeda_br(valor_decimal):
    return f"R$ {valor_decimal:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def obter_valor_diaria_individual(prestacao):
    """
    Retorna o valor individual de diária para a prestação/servidor.

    Prioridade:
    1) valor individual salvo explicitamente na própria prestação (dados_db);
    2) rateio do total do ofício pelo total de viajantes vinculados.
    """
    dados_db = prestacao.dados_db or {}
    valor_explicito = _parse_decimal(
        dados_db.get("valor_diaria_servidor")
        or dados_db.get("valor_diaria_individual")
        or dados_db.get("valor_saque")
    )
    if valor_explicito is not None:
        return valor_explicito

    total_oficio = _parse_decimal(getattr(prestacao.oficio, "valor_diarias", None))
    total_servidores = prestacao.oficio.viajantes.count() or 0
    if total_oficio is not None and total_servidores > 0:
        try:
            return (total_oficio / Decimal(total_servidores)).quantize(Decimal("0.01"))
        except (InvalidOperation, ZeroDivisionError):
            return None
    return None


def _dados_servidor(prestacao):
    servidor = prestacao.servidor
    if servidor:
        return {
            "nome_servidor": (servidor.nome or "").strip() or prestacao.nome_servidor,
            "rg_servidor": getattr(servidor, "rg_formatado", "") or prestacao.rg_servidor,
            "cpf_servidor": getattr(servidor, "cpf_formatado", "") or prestacao.cpf_servidor,
            "cargo_servidor": (getattr(servidor.cargo, "nome", "") or "").strip() or prestacao.cargo_servidor,
        }
    return {
        "nome_servidor": prestacao.nome_servidor,
        "rg_servidor": prestacao.rg_servidor,
        "cpf_servidor": prestacao.cpf_servidor,
        "cargo_servidor": prestacao.cargo_servidor,
    }


def _diaria_auto(prestacao):
    valor_individual = obter_valor_diaria_individual(prestacao)
    if valor_individual is not None:
        return _formatar_moeda_br(valor_individual)
    return "Nao informado"


def obter_ou_criar_rt(prestacao, usuario=None):
    dados_servidor = _dados_servidor(prestacao)
    defaults = {
        "oficio": prestacao.oficio,
        "servidor": prestacao.servidor,
        **dados_servidor,
        "diaria": _diaria_auto(prestacao),
        "translado": "Nao houve",
        "passagem": "Nao houve",
        "motivo": (prestacao.oficio.motivo or "").strip(),
        "atividade": "",
        "conclusao": "Todas as atividades previstas foram desenvolvidas.",
        "medidas": "Nada a acrescentar.",
        "informacoes_complementares": "",
    }
    rt, created = RelatorioTecnicoPrestacao.objects.get_or_create(prestacao=prestacao, defaults=defaults)
    if not created:
        changed = False
        for key, value in dados_servidor.items():
            if getattr(rt, key) != value:
                setattr(rt, key, value)
                changed = True
        diaria = _diaria_auto(prestacao)
        if rt.diaria != diaria:
            rt.diaria = diaria
            changed = True
        motivo = (prestacao.oficio.motivo or "").strip()
        if rt.motivo != motivo:
            rt.motivo = motivo
            changed = True
        if changed:
            rt.save()
    if created:
        prestacao.status_rt = PrestacaoConta.STATUS_RT_RASCUNHO
        prestacao.rt_atualizado_em = timezone.now()
        prestacao.save(update_fields=["status_rt", "rt_atualizado_em", "updated_at"])
    return rt


def montar_contexto_rt(rt):
    config = ConfiguracaoSistema.get_singleton()
    oficio = rt.oficio
    info_complementares = (rt.informacoes_complementares or "").strip() or INFO_COMPLEMENTARES_FALLBACK
    return {
        "divisao": _valor_padrao(config.divisao, "DEPARTAMENTO DA POLICIA CIVIL"),
        "unidade_cabecalho": _valor_padrao(config.unidade, "ASSESSORIA DE COMUNICACAO SOCIAL"),
        "oficio": _valor_padrao(oficio.numero_formatado, f"#{oficio.pk}"),
        "assunto_oficio": (oficio.motivo or "").strip(),
        "sede": _valor_padrao(config.sede, "Curitiba"),
        "data_atual_extenso": _formata_data_extenso_pt(timezone.localdate()),
        "nome_servidor": rt.nome_servidor,
        "rg_servidor": rt.rg_servidor,
        "diaria": _valor_padrao(rt.diaria, "Nao informado"),
        "translado": _valor_padrao(rt.translado, "Nao houve"),
        "passagem": _valor_padrao(rt.passagem, "Nao houve"),
        "motivo": (oficio.motivo or "").strip() or rt.motivo,
        "atividade": _valor_padrao(rt.atividade, ""),
        "conclusao": _valor_padrao(rt.conclusao, ""),
        "medidas": _valor_padrao(rt.medidas, ""),
        "informacoes_complementares": info_complementares,
        "info_complementares": info_complementares,
        "unidade_rodape": _valor_padrao(config.unidade, "ASSESSORIA DE COMUNICACAO SOCIAL"),
        "endereco": ", ".join(
            [p for p in [config.logradouro, config.numero, config.bairro, config.cidade_endereco, config.uf] if p]
        ),
        "telefone": _valor_padrao(config.telefone_formatado, "Nao informado"),
        "email": _valor_padrao(config.email, "nao informado"),
    }


def validar_placeholders_rt(caminho_docx_gerado):
    from docx import Document

    doc = Document(str(caminho_docx_gerado))
    leftovers = sorted(extract_placeholders_from_doc(doc))
    if leftovers:
        raise ValidationError(f"Ainda existem placeholders sem substituir: {', '.join(leftovers)}")


def _validar_placeholders_bytes(docx_bytes):
    from docx import Document

    doc = Document(BytesIO(docx_bytes))
    leftovers = sorted(extract_placeholders_from_doc(doc))
    if leftovers:
        raise ValidationError(f"Ainda existem placeholders sem substituir: {', '.join(leftovers)}")


def gerar_docx_rt(rt, usuario=None):
    if not RT_TEMPLATE_PATH.exists():
        raise ValidationError(f"Template DOCX nao encontrado em: {RT_TEMPLATE_PATH}")
    contexto = montar_contexto_rt(rt)
    docx_bytes = render_docx_template_bytes(RT_TEMPLATE_PATH, contexto)
    _validar_placeholders_bytes(docx_bytes)
    if PLACEHOLDER_PATTERN.search(docx_bytes.decode("latin-1", errors="ignore")):
        raise ValidationError("Documento DOCX gerado ainda contem placeholders brutos.")

    nome_servidor = slugify(rt.nome_servidor or "servidor")
    nome = f"relatorio-tecnico-oficio-{rt.oficio.pk}-{nome_servidor}.docx"
    rt.arquivo_docx.save(nome, ContentFile(docx_bytes), save=False)
    rt.marcar_como_gerado(usuario=usuario)
    rt.save(update_fields=["arquivo_docx", "updated_at"])
    return docx_bytes, nome


def gerar_pdf_rt(rt, usuario=None):
    if not rt.arquivo_docx:
        docx_bytes, _ = gerar_docx_rt(rt, usuario=usuario)
    else:
        rt.arquivo_docx.open("rb")
        try:
            docx_bytes = rt.arquivo_docx.read()
        finally:
            rt.arquivo_docx.close()
    try:
        pdf_bytes = convert_docx_bytes_to_pdf_bytes(docx_bytes)
    except (DocumentRendererUnavailable, DocumentValidationError) as exc:
        raise ValidationError(f"Conversao para PDF indisponivel neste ambiente: {exc}")
    nome_servidor = slugify(rt.nome_servidor or "servidor")
    nome = f"relatorio-tecnico-oficio-{rt.oficio.pk}-{nome_servidor}.pdf"
    rt.arquivo_pdf.save(nome, ContentFile(pdf_bytes), save=False)
    rt.marcar_como_gerado(usuario=usuario)
    rt.save(update_fields=["arquivo_pdf", "updated_at"])
    return pdf_bytes, nome
